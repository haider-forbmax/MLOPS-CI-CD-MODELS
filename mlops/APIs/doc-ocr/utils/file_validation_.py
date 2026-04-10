import os
from enum import Enum
from typing import Callable, Dict

from fastapi import HTTPException, UploadFile
from PIL import Image
from pypdf import PdfReader
from config import Config


class FileCategory(Enum):
    IMAGE = "image"
    PDF = "pdf"
    EXCEL = "excel"
    POWERPOINT = "powerpoint"
    DOCX = "docx"
    ODT = "odt"
    EPUB = "epub"
    DOCUMENT = "document"
    UNSUPPORTED = "unsupported"


def validate_uploaded_file(file: UploadFile, config: Config) -> None:
    """Validate uploaded file based on size, type, and structure."""
    file.file.seek(0, os.SEEK_END)
    file_size_mb = file.file.tell() / (1024 * 1024)
    file.file.seek(0)

    if file_size_mb > config.MAX_FILE_SIZE_MB:
        raise HTTPException(
            status_code=413,
            detail=f"File too large. Max {config.MAX_FILE_SIZE_MB} MB allowed."
        )

    _, ext = os.path.splitext(file.filename.lower())
    category = _get_file_category(ext, config)

    validators: Dict[FileCategory, Callable] = {
        FileCategory.IMAGE: lambda: _validate_image(file, config),
        FileCategory.PDF: lambda: _validate_pdf(file, config),
        FileCategory.EXCEL: lambda: _validate_excel(file, config),
        FileCategory.POWERPOINT: lambda: _validate_powerpoint(file, config),
        FileCategory.DOCX: lambda: _validate_docx(file, config),
        FileCategory.ODT: lambda: _validate_odt(file, config),
        FileCategory.EPUB: lambda: _validate_epub(file, config),
        FileCategory.DOCUMENT: lambda: _validate_document(file, config),
    }

    if category == FileCategory.UNSUPPORTED:
        raise HTTPException(415, detail=f"Unsupported file format: {ext}")

    try:
        validator = validators.get(category)
        if validator:
            validator()
    finally:
        file.file.seek(0)


def _get_file_category(ext: str, config: Config) -> FileCategory:
    ext_clean = ext.lstrip('.')
    if ext_clean in [e.lstrip('.') for e in config.IMAGES_TYPE]:
        return FileCategory.IMAGE
    elif ext_clean == 'pdf':
        return FileCategory.PDF
    elif ext_clean in [e.lstrip('.') for e in config.EXCEL_SUPPORTED_FORMATS]:
        return FileCategory.EXCEL
    elif ext_clean in [e.lstrip('.') for e in config.POWERPOINT_SUPPORTED_EXTENSIONS]:
        return FileCategory.POWERPOINT
    elif ext_clean == 'docx':
        return FileCategory.DOCX
    elif ext_clean == 'odt':
        return FileCategory.ODT
    elif ext_clean == 'epub':
        return FileCategory.EPUB
    elif ext_clean in [e.lstrip('.') for e in config.DOC_SUPPORTED_FORMATS]:
        return FileCategory.DOCUMENT
    return FileCategory.UNSUPPORTED


def _validate_image(file: UploadFile, config: Config) -> None:
    try:
        with Image.open(file.file) as image:
            if image.width > config.MAX_IMAGE_WIDTH or image.height > config.MAX_IMAGE_HEIGHT:
                raise HTTPException(
                    413,
                    detail=f"Image too large: {image.width}x{image.height}, "
                           f"max {config.MAX_IMAGE_WIDTH}x{config.MAX_IMAGE_HEIGHT} allowed."
                )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(400, detail=f"Invalid image file: {str(e)}")


def _validate_pdf(file: UploadFile, config: Config) -> None:
    try:
        reader = PdfReader(file.file, strict=False)
        if len(reader.pages) > config.MAX_PDF_PAGES:
            raise HTTPException(
                413, detail=f"PDF too long. Max {config.MAX_PDF_PAGES} pages allowed."
            )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(400, detail=f"Invalid PDF file: {str(e)}")


def _validate_excel(file: UploadFile, config: Config) -> None:
    import openpyxl
    wb = None
    try:
        wb = openpyxl.load_workbook(file.file, read_only=True, data_only=True, keep_links=False)
        if len(wb.sheetnames) > config.MAX_EXCEL_SHEETS:
            raise HTTPException(
                413, detail=f"Too many sheets. Max {config.MAX_EXCEL_SHEETS} allowed."
            )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(400, detail=f"Invalid Excel file: {str(e)}")
    finally:
        if wb:
            wb.close()


def _validate_powerpoint(file: UploadFile, config: Config) -> None:
    import tempfile
    from pptx import Presentation

    with tempfile.NamedTemporaryFile(suffix=".pptx", delete=True) as tmp:
        tmp.write(file.file.read())
        tmp.flush()
        file.file.seek(0)
        prs = Presentation(tmp.name)
        if len(prs.slides) > config.MAX_POWERPOINT_SLIDES:
            raise HTTPException(
                413, detail=f"Too many slides. Max {config.MAX_POWERPOINT_SLIDES} allowed."
            )


def _validate_docx(file: UploadFile, config: Config) -> None:
    from docx import Document
    try:
        doc = Document(file.file)
        if len(doc.paragraphs) > config.MAX_DOCX_PARAGRAPHS:
            raise HTTPException(
                413, detail=f"Too many paragraphs. Max {config.MAX_DOCX_PARAGRAPHS} allowed."
            )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(400, detail=f"Invalid DOCX file: {str(e)}")


def _validate_odt(file: UploadFile, config: Config) -> None:
    from odf import opendocument, text
    try:
        doc = opendocument.load(file.file)
        paragraphs = doc.getElementsByType(text.P)
        if len(paragraphs) > config.MAX_ODT_PARAGRAPHS:
            raise HTTPException(
                413, detail=f"Too many paragraphs. Max {config.MAX_ODT_PARAGRAPHS} allowed."
            )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(400, detail=f"Invalid ODT file: {str(e)}")


def _validate_epub(file: UploadFile, config: Config) -> None:
    import zipfile
    try:
        with zipfile.ZipFile(file.file, 'r') as epub:
            if len(epub.namelist()) > config.MAX_EPUB_FILES:
                raise HTTPException(
                    413, detail=f"EPUB too complex. Max {config.MAX_EPUB_FILES} files allowed."
                )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(400, detail=f"Invalid EPUB file: {str(e)}")


def _validate_document(file: UploadFile, config: Config) -> None:
    # No extra checks needed — file size already validated
    pass



def _get_natural_text_or_empty(self, ocr_text):
    import json

    try:
        data = json.loads(ocr_text)
    except Exception:
        return ""

    if "natural_text" not in data:
        return ""

    if data["natural_text"] is None:
        return ""

    return str(data["natural_text"])
