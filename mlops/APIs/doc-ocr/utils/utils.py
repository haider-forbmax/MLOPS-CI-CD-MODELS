import os
import socket
import shutil
import requests
import openpyxl
from PIL import Image
from config import Config
from PyPDF2 import PdfReader
from pptx import Presentation
from utils.file_converter_utils import *
from langdetect import detect_langs
from fastapi import HTTPException, UploadFile, Header


Config = Config()
base_url = f"http://{Config.URDU_API_BASE_URL}:3008"


def create_temp_folders(base_folder, subfolders):
    os.makedirs(base_folder, exist_ok=True)
    temp_paths = {}
    for subfolder in subfolders:
        subfolder_path = os.path.join(base_folder, subfolder)
        os.makedirs(subfolder_path, exist_ok=True)
        temp_paths[subfolder] = subfolder_path
    return temp_paths


def validate_api_key(api_key: str = Header(None, alias="X-API-Key")):
    keys = Config.USER_API_KEYS
    if api_key not in keys.values():
        raise HTTPException(status_code=401, detail="Invalid API key")
    return api_key


def prepare_upload_dir(upload_dir="uploads"):
    if os.path.exists(upload_dir):
        shutil.rmtree(upload_dir)
    os.makedirs(upload_dir, exist_ok=True)


def get_ip_address():
    try:
        host_name = socket.gethostname()
        ip_address = socket.gethostbyname(host_name)
        return ip_address
    except socket.error as e:
        print(f"Error: {e}")
        return None


def separate_urdu_english(result):
    english_text = ""
    urdu_text = ""
    modified_result = []
    for item in result:
        text = item[1]
        urdu_found = False
        try:
            if text:
                detected_languages = detect_langs(text)
                for lang in detected_languages:
                    if lang.lang in ("ur", "fa"):
                        urdu_found = True
                        break
                    else:
                        urdu_found = False
                if urdu_found:
                    urdu_text += text + " "
                else:
                    english_text += text + " "
                    continue
            modified_result.append(item)
        except:
            continue
    return english_text.strip(), urdu_text.strip(), modified_result


def read_text_file(file_path):
    try:
        with open(file_path, 'r') as file:
            return file.read()
    except FileNotFoundError:
        return f"File not found: {file_path}"
    except Exception as e:
        return f"An error occurred: {str(e)}"


def remove_null_and_empty(data):
    if isinstance(data, list):
        return [remove_null_and_empty(item) for item in data if item not in (None, "")]
    elif isinstance(data, dict):
        return {key: remove_null_and_empty(value) for key, value in data.items() if value not in (None, "")}
    else:
        return data


def log_gpu_memory():
    print("GPU logging disabled: this service is configured for CPU-only OCR.")


def validate_file_limits(file: UploadFile, config) -> None:
    file.file.seek(0, os.SEEK_END)
    file_size_mb = file.file.tell() / (1024 * 1024)
    file.file.seek(0)

    if file_size_mb > config.MAX_FILE_SIZE_MB:
        raise HTTPException(
            status_code=413,
            detail=f"File too large. Max {config.MAX_FILE_SIZE_MB} MB allowed."
        )

    _, ext = os.path.splitext(file.filename.lower())
    ext = file.filename.lower().split('.')[-1]

    if ext in config.IMAGES_TYPE:
        try:
            image = Image.open(file.file)
            if image.width > config.MAX_IMAGE_WIDTH or image.height > config.MAX_IMAGE_HEIGHT:
                raise HTTPException(
                    status_code=413,
                    detail=f"Image resolution too high. Max {config.MAX_IMAGE_WIDTH}x{config.MAX_IMAGE_HEIGHT} allowed."
                )
        except Exception:
            raise HTTPException(status_code=400, detail="Invalid image file")
        finally:
            file.file.seek(0)

    elif ext in config.DOC_SUPPORTED_FORMATS and ext == ".pdf":
        try:
            reader = PdfReader(file.file)
            if len(reader.pages) > config.MAX_PDF_PAGES:
                raise HTTPException(
                    status_code=413,
                    detail=f"PDF too long. Max {config.MAX_PDF_PAGES} pages allowed."
                )
        except Exception:
            raise HTTPException(status_code=400, detail="Invalid PDF file")
        finally:
            file.file.seek(0)

    # elif ext in config.EXCEL_SUPPORTED_FORMATS:
    #     try:
    #         wb = openpyxl.load_workbook(file.file, read_only=True)
    #         if len(wb.sheetnames) > config.MAX_EXCEL_SHEETS:
    #             raise HTTPException(
    #                 status_code=413,
    #                 detail=f"Excel file has too many sheets. Max {config.MAX_EXCEL_SHEETS} allowed."
    #             )
    #     except Exception as e:
    #         raise HTTPException(status_code=400, detail="Invalid Excel file")
    #     finally:
    #         file.file.seek(0)

    elif ext in config.POWERPOINT_SUPPORTED_EXTENSIONS:
        try:
            prs = Presentation(file.file)
            if len(prs.slides) > config.MAX_POWERPOINT_SLIDES:
                raise HTTPException(
                    status_code=413,
                    detail=f"PowerPoint file has too many slides. Max {config.MAX_POWERPOINT_SLIDES} allowed."
                )
        except Exception:
            raise HTTPException(status_code=400, detail="Invalid PowerPoint file")
        finally:
            file.file.seek(0)


    elif ext in config.DOC_SUPPORTED_FORMATS or ext == ".txt" or ext in config.EXCEL_SUPPORTED_FORMATS:
        # Already validated by size limit — no heavy structure
        pass

    else:
        raise HTTPException(
            status_code=415,
            detail=f"Unsupported file format: {ext}"
        )




def urdu_ocr(file_path):
    try:
        with open(file_path, 'rb') as file_:
            files = {'file': file_}
            headers = {'api_key': Config.API_KEY}
            response = requests.post(f"{base_url}/urdu_ocr/", files=files, headers=headers)

        if response.status_code == 200:
            return response.json()
        return f"Error: {response.status_code} - {response.text}"

    except requests.exceptions.RequestException as e:
        return f"Request error: {e}"
    except Exception as e:
        return f"General error: {str(e)}"


def initialize_ocr_reader():
    try:
        import easyocr
    except ImportError as exc:
        raise RuntimeError(
            "Legacy EasyOCR pipeline is not installed in this CPU-optimized build."
        ) from exc

    return easyocr.Reader(['ur', 'en'], gpu=False)


def make_pdf(file_path, extension, output_folder):
    ext = extension.lower()

    if ext == "doc":
        return doc_to_pdf(file_path, output_folder)
    elif ext in Config.EXCEL_SUPPORTED_FORMATS:
        return excel_to_pdf(file_path, output_folder)
    elif ext in Config.DOC_SUPPORTED_FORMATS:
        return convert_to_pdf(file_path, output_folder)
    elif ext in Config.POWERPOINT_SUPPORTED_EXTENSIONS:
        return ppt_to_pdf(file_path, output_folder)
    else:
        return libreoffice_convert_to_pdf(file_path, output_folder)
